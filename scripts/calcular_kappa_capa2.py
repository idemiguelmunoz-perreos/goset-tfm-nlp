"""Acuerdo inter-evaluador de la Capa 2 (validación experta del recomendador).

Sobre una escala Likert ordinal 1-4, el kappa ponderado de Cohen se deprime ante
marginales concentradas (paradoja de Feinstein y Cicchetti, 1990): con acuerdo
observado alto pero puntuaciones agolpadas en 3-4, kappa tiende a cero. Por eso
este módulo reporta, además del kappa ponderado cuadrático, el coeficiente AC2 de
Gwet (2008), robusto a esa paradoja, junto con acuerdo observado ponderado,
acuerdo dentro de +-1, porcentaje de recomendaciones apropiadas (>=3) y el sesgo
sistemático entre evaluadoras.

Determinista: no usa aleatoriedad. Uso:
    PYTHONPATH=src python scripts/calcular_kappa_capa2.py docs/kit_capa2_REAL_96.xlsx
"""

from __future__ import annotations

import sys
from collections import Counter
from itertools import combinations

from openpyxl import load_workbook

RATERS: dict[str, int] = {"Vet": 4, "Peluquera": 5, "Isabela": 6}  # columnas 1-indexadas
K: int = 4  # categorías de la escala Likert
CATS: tuple[int, ...] = tuple(range(1, K + 1))


def _peso_acuerdo(i: int, j: int) -> float:
    """Peso de acuerdo cuadrático: 1 en la diagonal, 0 en el extremo opuesto."""
    return 1.0 - (i - j) ** 2 / (K - 1) ** 2


def kappa_ponderado(pares: list[tuple[int, int]]) -> float:
    """Cohen's kappa con pesos cuadráticos (desacuerdo) para escalas ordinales."""
    n = len(pares)
    if n == 0:
        return float("nan")
    obs = Counter(pares)
    ma = {c: sum(1 for a, _ in pares if a == c) / n for c in CATS}
    mb = {c: sum(1 for _, b in pares if b == c) / n for c in CATS}
    dis = lambda i, j: (i - j) ** 2 / (K - 1) ** 2  # noqa: E731
    po = sum(dis(i, j) * obs[(i, j)] / n for i in CATS for j in CATS)
    pe = sum(dis(i, j) * ma[i] * mb[j] for i in CATS for j in CATS)
    return 1.0 if pe == 0 else 1.0 - po / pe


def gwet_ac2(pares: list[tuple[int, int]]) -> float:
    """AC2 de Gwet con pesos cuadráticos; robusto a la paradoja del kappa."""
    n = len(pares)
    if n == 0:
        return float("nan")
    obs = Counter(pares)
    ca = Counter(a for a, _ in pares)
    cb = Counter(b for _, b in pares)
    pi = {c: ((ca[c] / n) + (cb[c] / n)) / 2 for c in CATS}
    tw = sum(_peso_acuerdo(i, j) for i in CATS for j in CATS)
    pa = sum(_peso_acuerdo(i, j) * obs[(i, j)] / n for i in CATS for j in CATS)
    pe = (tw / (K * (K - 1))) * sum(pi[c] * (1 - pi[c]) for c in CATS)
    return (pa - pe) / (1 - pe) if pe != 1 else 1.0


def acuerdo_observado(pares: list[tuple[int, int]]) -> float:
    """Acuerdo observado ponderado cuadrático (Pa)."""
    n = len(pares)
    return sum(_peso_acuerdo(a, b) for a, b in pares) / n if n else float("nan")


def dentro_de_uno(pares: list[tuple[int, int]]) -> float:
    """Proporción de pares que difieren en +-1 punto o menos."""
    n = len(pares)
    return sum(1 for a, b in pares if abs(a - b) <= 1) / n if n else float("nan")


def _leer(path: str) -> dict[str, list[int]]:
    ws = load_workbook(path)["Evaluacion"]
    puntuaciones: dict[str, list[int]] = {r: [] for r in RATERS}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or str(row[0]).strip().upper() == "EJEMPLO":
            continue
        for r, col in RATERS.items():
            v = row[col - 1]
            if isinstance(v, (int, float)) and v:
                puntuaciones[r].append(int(v))
    return puntuaciones


def main(path: str) -> None:
    p = _leer(path)
    n = len(next(iter(p.values())))
    print(f"Ítems evaluados: {n}\n")

    print("Por evaluadora:")
    for r in RATERS:
        v = p[r]
        media = sum(v) / len(v)
        apropiado = sum(1 for x in v if x >= 3) / len(v)
        print(f"  {r}: media Likert={media:.2f}  %apropiado(>=3)={apropiado:.1%}")

    print("\nAcuerdo entre pares (n={}):".format(n))
    print(f"  {'par':22} {'kappa_w':>8} {'AC2_w':>7} {'Pa_w':>6} {'±1':>6}")
    for a, b in combinations(RATERS, 2):
        pares = list(zip(p[a], p[b]))
        print(
            f"  {a + ' vs ' + b:22} {kappa_ponderado(pares):8.3f} "
            f"{gwet_ac2(pares):7.3f} {acuerdo_observado(pares):6.3f} "
            f"{dentro_de_uno(pares):6.1%}"
        )

    print("\nSesgo sistemático (diferencia de medias emparejadas a - b):")
    for a, b in combinations(RATERS, 2):
        d = [x - y for x, y in zip(p[a], p[b])]
        print(f"  {a} - {b}: {sum(d) / len(d):+.3f}")

    allv = [x for r in RATERS for x in p[r]]
    print(
        f"\nAgregado: media global={sum(allv) / len(allv):.2f}  "
        f"%apropiado(>=3)={sum(1 for x in allv if x >= 3) / len(allv):.1%}"
    )


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "docs/kit_capa2_REAL_96.xlsx")

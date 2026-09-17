"""Perfiles reales completos (formulario + cartilla) → recomendaciones (Componente B).

Une las respuestas del cuestionario (features de comportamiento) con los datos
clínicos de la cartilla por 'ID cartilla', y ejecuta el recomendador. Anonimiza:
no conserva nombre ni contacto del dueño.
"""
from __future__ import annotations
import json, sys, pathlib
from openpyxl import load_workbook
from goset_recomendador.engine import recomendar

MAP={"5. Raza o mezcla":"Raza o mezcla","6. Edad aproximada":"Edad (años)","7. Sexo":"Sexo",
 "8. ¿Está castrado/esterilizado?":"Castrado/esterilizado","10. Peso aproximado (kg)":"Peso (kg)",
 "11. Tipo de pelo":"Tipo de manto","12. ¿Se le enreda o apelmaza el pelo con facilidad?":"¿Se enreda/apelmaza?",
 "13. Orejas":"Orejas","14. ¿Problemas de piel, alergias o dermatitis?":"Problemas de piel/dermatitis",
 "16. ¿Tiene miedo al agua o al baño?":"Miedo al agua","17. Nivel de actividad":"Nivel de actividad",
 "21. Con otros perros":"Con otros perros","22. Con personas desconocidas":"Con personas",
 "23. ¿Ansiedad cuando se queda solo?":"Ansiedad por separación",
 "24. ¿Miedo a ruidos, tormentas o ansiedad general?":"Miedo a ruidos/tormentas",
 "25. ¿Protege comida/juguetes o ha mostrado agresividad?":"Protección de recursos/agresividad",
 "26. Nivel de educación / obediencia":"Nivel de educación/obediencia",
 "27. Estado corporal percibido (mirándolo y tocándole las costillas)":"BCS (estado corporal)",
 "28. ¿Movilidad reducida, artrosis o problemas articulares?":"Movilidad/artrosis",
 "29. Si es mayor: ¿signos de desorientación o confusión?":"Disfunción cognitiva (senior)",
 "30. Tipo de dieta":"Tipo de dieta","2. Nombre del perro":"Nombre del perro"}
CART_MAP={"rabia_valida_hasta":"Rabia (válida hasta)","polivalente_valida_hasta":"Polivalente (válida hasta)",
 "despar_fecha":"Desparasitación (fecha)","despar_tipo":"Desparasitación (tipo)"}

def rows(ws):
    hdr=[c.value for c in ws[1]]
    return hdr,[dict(zip(hdr,r)) for r in ws.iter_rows(min_row=2,values_only=True) if any(v not in (None,"") for v in r)]

def main(form_path, cart_path):
    _,frows=rows(load_workbook(form_path)["Respuestas"])
    _,crows=rows(load_workbook(cart_path)["GoldenSet"])
    clin={c["doc_id"]:c for c in crows if c.get("doc_id")}
    perfiles=[]
    for i,fr in enumerate(frows):
        idc=fr.get("ID cartilla (trazabilidad)")
        p={"id":f"real_{i+1:03d}"}
        for src,dst in MAP.items():
            if fr.get(src) not in (None,""): p[dst]=fr[src]
        cc=clin.get(idc,{})
        for src,dst in CART_MAP.items():
            if cc.get(src) not in (None,""): p[dst]=cc[src]
        perfiles.append(p)
    outs=[recomendar(p) for p in perfiles]
    root=pathlib.Path("tests/fixtures")
    json.dump({"_nota":"Perfiles REALES (formulario+cartilla), identidad del dueño eliminada","n":len(perfiles),"perfiles":perfiles},
              open(root/"perfiles_reales.json","w"),ensure_ascii=False,indent=2)
    json.dump({"n":len(outs),"recomendaciones":outs},open(root/"recomendaciones_reales.json","w"),ensure_ascii=False,indent=2)
    from collections import Counter
    top=Counter(o["servicio_top"] for o in outs)
    print(f"Perfiles reales procesados: {len(perfiles)}")
    print("Servicio prioritario (top-1):", dict(top.most_common()))
    print("needs_human_review:", sum(1 for o in outs if o["flags"]["needs_human_review"]))
    print("no elegibles a grupo:", sum(1 for o in outs if not o["flags"]["elegible_grupo"]))

if __name__=="__main__":
    main(sys.argv[1], sys.argv[2])

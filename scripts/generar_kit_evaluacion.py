"""Genera el kit de evaluación Capa 2 (validación experta) en un Excel.

Toma una muestra de perfiles, corre el recomendador, y produce una hoja donde
la veterinaria, la peluquera e Isabela puntúan cada recomendación (Likert 1-4)
de forma independiente. El acuerdo se calcula luego con scripts/calcular_kappa_capa2.py.

Uso:
    PYTHONPATH=src python scripts/generar_kit_evaluacion.py [n]
"""

from __future__ import annotations

import json
import pathlib
import random
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from goset_recomendador.engine import recomendar
from goset_recomendador.normalize import normalize

FIX = pathlib.Path("tests/fixtures/perfiles_sinteticos.json")
OUT = pathlib.Path("docs/kit_evaluacion_capa2.xlsx")
FNT = "Arial"


def _resumen(p: dict) -> str:
    f = normalize(p)
    partes = [p.get("Nombre del perro") or p.get("id"), p.get("Raza o mezcla"),
              f"{f['edad']}a" if f["edad"] else None, p.get("Sexo")]
    if f["hembra_entera"]:
        partes.append("entera")
    if f["manto"]:
        partes.append(f"manto {f['manto']}")
    señas = []
    if f["reactivo_perros"] or f["reactivo_personas"]:
        señas.append("reactivo")
    if f["ansiedad_sep"]:
        señas.append("ansiedad separación")
    if f["bcs_sobrepeso"]:
        señas.append("sobrepeso")
    if f["movilidad_reducida"]:
        señas.append("movilidad reducida")
    base = " · ".join(str(x) for x in partes if x)
    return base + (f" · {', '.join(señas)}" if señas else "")


def _reco_texto(out: dict) -> str:
    lineas = []
    for r in out["ranking"][:3]:
        ids = ", ".join(rr["id"] for rr in r["reglas_activadas"])
        lineas.append(f"{r['servicio']} (prio {r['prioridad']}; reglas {ids})")
    return " | ".join(lineas)


def main() -> None:
    n = int(sys.argv[1]) if len(sys.argv) > 1 else 20
    perfiles = json.loads(FIX.read_text("utf-8"))["perfiles"]
    random.seed(42)
    muestra = random.sample(perfiles, min(n, len(perfiles)))

    wb = Workbook()
    hf = PatternFill("solid", fgColor="1F3864")
    hff = Font(name=FNT, bold=True, color="FFFFFF", size=10)
    yellow = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="D9D9D9")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)

    ins = wb.active
    ins.title = "Instrucciones"
    filas = [
        ["Kit de evaluación — Capa 2 (validación experta del recomendador)"],
        [""],
        ["Cómo se rellena:"],
        ["1. Cada evaluadora puntúa DE FORMA INDEPENDIENTE, sin ver las notas de las otras."],
        ["2. Escala Likert (columna por evaluadora):"],
        ["   1 = Recomendación inapropiada / incorrecta"],
        ["   2 = Dudosa / discutible"],
        ["   3 = Apropiada con reservas"],
        ["   4 = Totalmente apropiada"],
        ["3. Rellenad SOLO las celdas amarillas (Vet, Peluquera, Isabela) y comentarios."],
        ["4. La fila marcada EJEMPLO no se puntúa: es solo muestra de formato."],
        [""],
        ["Nota: muestra sobre perfiles sintéticos (desarrollo). Valida la solidez de la"],
        ["base de conocimiento, no el rendimiento sobre perros reales."],
    ]
    for r in filas:
        ins.append(r)
    ins["A1"].font = Font(name=FNT, bold=True, size=12, color="1F3864")
    ins.column_dimensions["A"].width = 90

    ev = wb.create_sheet("Evaluacion")
    headers = ["ID", "Perro (resumen)", "Recomendación del sistema (top-3)",
               "Vet (1-4)", "Peluquera (1-4)", "Isabela (1-4)", "Comentarios"]
    ev.append(headers)
    for c in ev[1]:
        c.fill = hf; c.font = hff; c.border = bd
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    # fila EJEMPLO
    ev.append(["EJEMPLO", "Rex · Caniche · 3a · macho · manto rizado",
               "Peluqueria (prio Media; reglas PE-01, PE-07) | Salud preventiva (...)",
               4, 4, 3, "De acuerdo; quizá cada 5 semanas"])
    for c in ev[2]:
        c.font = Font(name=FNT, size=9, italic=True, color="808080")
        c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = bd

    for p in muestra:
        out = recomendar(p)
        ev.append([p.get("id"), _resumen(p), _reco_texto(out), "", "", "", ""])
        rr = ev.max_row
        for c in ev[rr]:
            c.font = Font(name=FNT, size=9); c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = bd
        for col in (4, 5, 6, 7):
            ev.cell(rr, col).fill = yellow

    W = {"A": 14, "B": 34, "C": 52, "D": 10, "E": 12, "F": 11, "G": 30}
    for k, v in W.items():
        ev.column_dimensions[k].width = v
    ev.freeze_panes = "A2"
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"Kit generado: {OUT} | perfiles evaluables: {len(muestra)}")


if __name__ == "__main__":
    main()
